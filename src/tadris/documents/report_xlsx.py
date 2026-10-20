"""A report as an Excel workbook: one sheet, real numbers and dates (sums and filters work in
Excel), the same rows and total as on screen."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from tadris.core.cells import excel_value
from tadris.domain.models import ReportTable

FORMATS = {
    "number": "#,##0",
    "money": '#,##0 "DA"',
    "percent": "0.0%",
    "date": "DD/MM/YYYY",
    "datetime": "DD/MM/YYYY HH:MM",
}
WIDTHS = {"text": 28, "number": 10, "money": 15, "percent": 10, "date": 12, "datetime": 17}
HEADER_ROW = 4
SHEET_TITLE_MAX = 31  # Excel's limit


def build_report_xlsx(report: ReportTable, school: str, path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = report.title[:SHEET_TITLE_MAX]
    sheet["A1"] = report.title
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = f"{school} · {report.period}" if school else report.period
    sheet["A2"].font = Font(italic=True, color="5A5868")

    shade = PatternFill("solid", fgColor="EFEDF5")
    for index, column in enumerate(report.columns, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=column.label)
        cell.font = Font(bold=True)
        cell.fill = shade
        cell.alignment = Alignment(horizontal="left" if column.kind == "text" else "right",
                                   wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS[column.kind]

    def write(row_index: int, values: tuple, bold: bool = False) -> None:
        for index, (column, value) in enumerate(zip(report.columns, values, strict=True),
                                                start=1):
            cell = sheet.cell(row=row_index, column=index, value=excel_value(column.kind, value))
            if column.kind in FORMATS:
                cell.number_format = FORMATS[column.kind]
            if bold:
                cell.font = Font(bold=True)
                cell.border = Border(top=Side(style="thin"))

    row_index = HEADER_ROW
    for values in report.rows:
        row_index += 1
        write(row_index, values)
    if report.total is not None:
        row_index += 1
        write(row_index, report.total, bold=True)
    if report.note:
        sheet.cell(row=row_index + 2, column=1, value=report.note).font = Font(
            italic=True, color="5A5868")
    sheet.freeze_panes = sheet.cell(row=HEADER_ROW + 1, column=1)
    if report.rows:
        last = get_column_letter(len(report.columns))
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{last}{HEADER_ROW + len(report.rows)}"
    book.save(path)
    return path
