"""Long lists come a page at a time: the students, the student list report, a student's
attendance, any report on screen; debts and attendance counts are computed by the database
over the whole history."""

from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import Qt

from tadris import config
from tadris.core import dates
from tadris.core.errors import ValidationError
from tadris.db.connection import transaction
from tadris.services import attendance, auth, debt, students
from tadris.ui.pages.reports import ReportsPage
from tadris.ui.pages.student_attendance_tab import StudentAttendanceTab
from tadris.ui.pages.students import StudentListView
from tadris.ui.theme import apply_theme, load_fonts
from tadris.ui.widgets import common
from tadris.ui.widgets.tables import PAGE_SIZE

TODAY = date(2026, 10, 18)


@pytest.fixture(autouse=True)
def theme(qapp, monkeypatch):
    load_fonts()
    apply_theme(qapp)
    monkeypatch.setattr(dates, "today", lambda: TODAY)


@pytest.fixture(autouse=True)
def documents(monkeypatch, tmp_path):
    opened = []
    monkeypatch.setattr(config, "documents_dir", lambda: tmp_path)
    monkeypatch.setattr(common, "open_document", lambda path: opened.append(path))
    return opened


@pytest.fixture
def director(conn):
    return auth.bootstrap_director(conn, "directeur", "password1", "École Ennour")


NAMES = ["Amrani", "Benali", "Cherif", "Djebbar", "Ferhat", "Haddad", "Kaci", "Zerrouki"]


@pytest.fixture
def many(conn):
    """250 students, names spread over the alphabet."""
    with transaction(conn):
        for n in range(250):
            conn.execute("INSERT INTO students(student_number, first_name, last_name, "
                         "birth_date) VALUES (?, ?, ?, ?)",
                         (f"2026-{n + 1:04d}", f"Prénom{n:03d}", NAMES[n % len(NAMES)],
                          f"20{10 + n % 9:02d}-01-01"))
    return 250


def scroll_until(qtbot, table, rows: int) -> None:
    bar = table.verticalScrollBar()

    def there() -> bool:
        bar.setValue(bar.maximum())
        return table.rowCount() >= rows

    qtbot.waitUntil(there, timeout=5000)


# ----- the students list -----

def test_students_search_reads_pages_in_any_order(conn, director, many):
    page1 = students.search(conn, director, limit=PAGE_SIZE)
    page2 = students.search(conn, director, limit=PAGE_SIZE, offset=PAGE_SIZE)
    everyone = students.search(conn, director)
    assert [s.id for s in page1 + page2] == [s.id for s in everyone[:200]]
    assert students.count(conn, director) == 250
    assert students.count(conn, director, "zerrouki") == 31
    last_first = students.search(conn, director, order="last_name", descending=True, limit=1)
    assert last_first[0].last_name == "Zerrouki"
    by_number = students.search(conn, director, order="number", limit=3)
    assert [s.student_number for s in by_number] == ["2026-0001", "2026-0002", "2026-0003"]
    with pytest.raises(ValidationError):
        students.search(conn, director, order="password_hash")  # only the list's columns


def test_the_students_page_loads_more_as_it_is_scrolled(qtbot, conn, director, many):
    view = StudentListView(conn, director)
    qtbot.addWidget(view)
    view.resize(1200, 700)
    view.show()
    assert view.table.rowCount() == PAGE_SIZE
    assert view.count.text().startswith("100 affichés sur 250 élève(s)")
    scroll_until(qtbot, view.table, 200)
    assert len(set(view.student_ids())) == view.table.rowCount()  # no student twice
    scroll_until(qtbot, view.table, 250)
    assert view.count.text() == "250 élève(s)"


def test_a_column_click_sorts_every_student(qtbot, conn, director, many):
    view = StudentListView(conn, director)
    qtbot.addWidget(view)
    view.show()
    view.table.horizontalHeader().setSortIndicator(1, Qt.SortOrder.DescendingOrder)
    # The first row is the last name of all 250, not of the 100 loaded.
    assert view.table.item(0, 1).text() == "Zerrouki"
    assert view.table.rowCount() == PAGE_SIZE


# ----- the student list report -----

def test_the_student_report_comes_a_page_at_a_time(qtbot, conn, director, many, documents):
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    page.resize(1300, 700)
    page.show()
    page.tabs.setCurrentIndex([*page.views].index("students"))
    view = page.views["students"]
    view.status.setCurrentIndex(view.status.findData(None))
    assert view.table.rowCount() == PAGE_SIZE + 1  # and the total row
    assert view.table.item(PAGE_SIZE, 0).text() == "250 élève(s)"  # for all of them
    scroll_until(qtbot, view.table, 251)
    assert view.table.item(250, 0).text() == "250 élève(s)"  # still at the bottom
    assert len(set(view.report.row_ids)) == 250
    with qtbot.waitSignal(view.exported, timeout=15000) as done:
        qtbot.mouseClick(view.excel, Qt.MouseButton.LeftButton)
    sheet = load_workbook(done.args[0]).active
    assert sheet.max_row - 4 - 1 - 2 == 250  # rows, then the total and the note


# ----- any report on screen -----

def test_a_long_report_is_shown_a_page_at_a_time(qtbot, conn, director):
    from tadris.ui.pages.reports import ReportView

    rows = tuple((f"Élève {n:03d}", n) for n in range(260))
    from tadris.domain.models import ReportColumn, ReportTable

    report = ReportTable(key="x", title="X", period="P",
                         columns=(ReportColumn("Élève"), ReportColumn("N", "number")),
                         rows=rows, total=("Total", sum(range(260))))
    view = ReportView(conn, director, lambda: report)
    qtbot.addWidget(view)
    view.resize(900, 600)
    view.show()
    view.refresh()
    assert view.table.rowCount() == PAGE_SIZE + 1
    scroll_until(qtbot, view.table, 261)
    assert view.rows_shown()[-1] == ["Total", f"{sum(range(260)):,}".replace(",", " ")]
    assert view.rows_shown()[259][0] == "Élève 259"


# ----- a student's attendance -----

@pytest.fixture
def marked(conn, director):
    """One student, 260 sessions of their group, every one marked."""
    with transaction(conn):
        for statement in """
        INSERT INTO school_years(label, start_date, end_date) VALUES('2026-2027','2024-09-01','2027-06-30');
        INSERT INTO levels(name) VALUES('3AM');
        INSERT INTO subjects(name) VALUES('Math');
        INSERT INTO rooms(name, capacity) VALUES('Salle A', 20);
        INSERT INTO students(student_number,first_name,last_name) VALUES('2026-9999','Amine','Benali');
        INSERT INTO teachers(full_name) VALUES('T');
        INSERT INTO registrations(student_id,school_year_id,level_id,fee_amount) VALUES(1,1,1,0);
        INSERT INTO class_groups(school_year_id,level_id,subject_id,name,price,capacity,start_date)
          VALUES(1,1,1,'G1',3000,10,'2024-09-01');
        INSERT INTO enrollments(student_id,group_id,registration_id,enrolled_on) VALUES(1,1,1,'2024-09-01');
        """.strip().split(";\n"):
            conn.execute(statement.strip().rstrip(";"))
        day = date(2024, 9, 2)
        for n in range(260):
            session = conn.execute(
                "INSERT INTO sessions(group_id, room_id, teacher_id, date, start_time, "
                "end_time, status, source) VALUES (1, 1, 1, ?, '17:00', '18:30', 'held', "
                "'manual')", (day.isoformat(),)).lastrowid
            conn.execute("INSERT INTO attendances(session_id, enrollment_id, status, "
                         "entered_by) VALUES (?, 1, ?, ?)",
                         (session, "absent" if n % 10 == 0 else "present", director.user_id))
            day += timedelta(days=3)
    return 1


def test_a_students_attendance_counts_all_and_lists_a_page(qtbot, conn, director, marked):
    summary = attendance.summary_for_student(conn, director, marked)
    assert (summary.present, summary.absent, summary.total) == (234, 26, 260)
    tab = StudentAttendanceTab(conn, director, marked)
    qtbot.addWidget(tab)
    tab.resize(900, 600)
    tab.show()
    assert "234" in tab.summary.text() and "26" in tab.summary.text()  # the whole history
    table = tab.list.table
    assert table.rowCount() == PAGE_SIZE
    first = attendance.for_student(conn, director, marked, limit=1)[0]
    assert table.item(0, 0).text() == dates.display_date(first.date)  # latest first
    scroll_until(qtbot, table, 260)
    assert table.rowCount() == 260


# ----- debts read what is still owed -----

def test_debts_leave_out_what_is_paid(conn, director):
    from test_reports import School

    class Patch:
        def setattr(self, obj, name, value):
            setattr(obj, name, value)

    saved = dates.today
    try:
        school = School(conn, director, Patch())
    finally:
        dates.today = saved
    from tadris.repositories import payments as repo

    everyone = repo.open_charges(conn)
    assert everyone and all(row["paid"] < row["amount"] for row in everyone)
    paid_in_full = conn.execute(
        "SELECT COUNT(*) FROM charges c WHERE c.status = 'open' AND c.amount <= "
        "(SELECT COALESCE(SUM(amount), 0) FROM payment_allocations pa WHERE pa.charge_id = c.id "
        "AND NOT EXISTS (SELECT 1 FROM payment_voids v WHERE v.payment_id = pa.payment_id))"
    ).fetchone()[0]
    assert paid_in_full > 0  # there are some, and they are not read
    for student in (school.amine, school.badis, school.chahinez):
        mine = [dict(r) for r in repo.open_charges(conn, student)]
        theirs = [dict(r) for r in everyone if r["student_id"] == student]
        assert mine == theirs  # one student's query and everyone's agree
    assert debt.every_debt(conn)[school.badis].owed_now == 7000
    assert school.amine not in debt.every_debt(conn)  # paid everything

