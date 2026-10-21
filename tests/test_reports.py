import json
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from tadris.core import cells, dates
from tadris.core.context import UserContext
from tadris.core.errors import PermissionDenied, ValidationError
from tadris.domain.models import ChargeDue, Debt, GroupInput, StudentInput, TeacherInput
from tadris.i18n.fr import t
from tadris.services import (
    auth,
    catalog,
    enrollments,
    groups,
    payments,
    payouts,
    printing,
    registrations,
    reports,
    school_years,
    students,
    subscriptions,
    teachers,
)

TODAY = date(2026, 10, 18)


def move_to(monkeypatch, day: date) -> None:
    monkeypatch.setattr(dates, "today", lambda: day)


@pytest.fixture
def director(conn):
    return auth.bootstrap_director(conn, "directeur", "password1", "École Ennour")


@pytest.fixture
def desk(conn, director):
    user_id = auth.create_user(conn, director, "accueil", "password2", "receptionist")
    return UserContext(user_id=user_id, username="accueil", role="receptionist")


class School:
    """Karim (60 %) teaches A, Maths, 6 000 DA; Samia (50 %) teaches B, Physique, 4 000 DA.

    01/09  Amine: fee 1 000, A from 01/09.  Badis: fee 1 000, B from 01/09.
    02/09  Amine pays 7 000 (fee + September).  05/09  Badis pays 2 000 (fee + 1 000).
    10/09  Badis pays 500, voided the same day.
    01/10  Amine and Badis renewed for October.
    02/10  Karim's September (3 600) paid; Samia's September (2 000) left as a draft.
    03/10  Amine pays 6 000 (October).
    05/10  Chahinez: fee 1 000, A from 05/10.  06/10  she pays 8 000 (1 000 left as credit).
    Today 18/10.

    By hand: September charged 12 000, settled 9 000, collected 9 000.
             October charged 17 000, settled 13 000, collected 14 000.
             Badis owes 3 000 from 01/09 (47 days) and 4 000 from 01/10 (17 days).
    """

    def __init__(self, conn, ctx, monkeypatch):
        self.conn, self.ctx = conn, ctx
        move_to(monkeypatch, date(2026, 9, 1))
        self.year = school_years.create(conn, ctx, "2026-2027", date(2026, 9, 1),
                                        date(2027, 6, 30))
        self.level = catalog.create(conn, ctx, "level", "4AM")
        self.karim = teachers.create(conn, ctx, TeacherInput("Karim Haddad"), 6000,
                                     date(2026, 9, 1))
        self.samia = teachers.create(conn, ctx, TeacherInput("Samia Rahmani"), 5000,
                                     date(2026, 9, 1))
        self.a = self.group("A", "Maths", 6000, self.karim)
        self.b = self.group("B", "Physique", 4000, self.samia)

        self.amine = self.pupil("Amine", "Benali", self.a)
        self.badis = self.pupil("Badis", "Cherif", self.b)
        move_to(monkeypatch, date(2026, 9, 2))
        payments.record(conn, ctx, self.amine, 7000, "cash")
        move_to(monkeypatch, date(2026, 9, 5))
        payments.record(conn, ctx, self.badis, 2000, "cash")
        move_to(monkeypatch, date(2026, 9, 10))
        voided = payments.record(conn, ctx, self.badis, 500, "cash")
        payments.void(conn, ctx, voided, "Erreur de saisie")

        move_to(monkeypatch, date(2026, 10, 1))
        for student in (self.amine, self.badis):
            [month] = subscriptions.for_student(conn, ctx, student)
            subscriptions.renew(conn, ctx, month.id)
        move_to(monkeypatch, date(2026, 10, 2))
        paid = payouts.build_draft(conn, ctx, self.karim, date(2026, 9, 30))
        payouts.approve(conn, ctx, paid)
        payouts.mark_paid(conn, ctx, paid, "cash")
        payouts.build_draft(conn, ctx, self.samia, date(2026, 9, 30))
        move_to(monkeypatch, date(2026, 10, 3))
        payments.record(conn, ctx, self.amine, 6000, "cash")

        move_to(monkeypatch, date(2026, 10, 5))
        self.chahinez = self.pupil("Chahinez", "Djebbar", self.a)
        move_to(monkeypatch, date(2026, 10, 6))
        payments.record(conn, ctx, self.chahinez, 8000, "transfer")
        move_to(monkeypatch, TODAY)

    def group(self, name, subject, price, teacher):
        subject_id = catalog.create(self.conn, self.ctx, "subject", subject)
        return groups.create(self.conn, self.ctx, GroupInput(
            self.year, self.level, subject_id, name, price, 20, date(2026, 9, 1)), teacher)

    def pupil(self, first, last, group_id) -> int:
        student = students.create(self.conn, self.ctx, StudentInput(first, last))
        registrations.register(self.conn, self.ctx, student, self.level, 1000)
        enrollments.enroll(self.conn, self.ctx, student, group_id, enrolled_on=dates.today())
        return student


@pytest.fixture
def school(conn, director, monkeypatch):
    return School(conn, director, monkeypatch)


def column(table, label):
    return [c.label for c in table.columns].index(label)


# ----- revenue -----

def test_revenue_by_month_matches_the_hand_count(conn, director, school):
    table = reports.revenue_by_month(conn, director, school.year)
    assert [c.label for c in table.columns] == [
        "Mois", "Facturé", "Réglé", "Reste à payer", "Encaissé"]
    assert len(table.rows) == 10  # September to June
    assert table.rows[0] == ("Septembre 2026", 12000, 9000, 3000, 9000)
    assert table.rows[1] == ("Octobre 2026", 17000, 13000, 4000, 14000)
    assert all(row[1:] == (0, 0, 0, 0) for row in table.rows[2:])
    assert table.rows[-1][0] == "Juin 2027"
    assert table.total == ("Total", 29000, 22000, 7000, 23000)
    assert table.period == "Année 2026-2027"


def test_revenue_by_group_over_the_year_and_one_month(conn, director, school):
    year = reports.revenue_by_group(conn, director, school.year)
    assert year.rows == (
        ("A · Maths 4AM", 3, 18000, 18000, 0),
        ("B · Physique 4AM", 2, 8000, 1000, 7000),
        ("Frais d'inscription", 3, 3000, 3000, 0),
    )
    assert year.total == ("Total", 8, 29000, 22000, 7000)

    october = reports.revenue_by_group(conn, director, school.year, date(2026, 10, 1))
    assert october.period == "Octobre 2026"
    assert october.rows == (
        ("A · Maths 4AM", 2, 12000, 12000, 0),
        ("B · Physique 4AM", 1, 4000, 0, 4000),
        ("Frais d'inscription", 1, 1000, 1000, 0),
    )
    assert october.total[2:4] == (17000, 13000)  # = October in the monthly report


def test_a_cancelled_month_is_not_charged(conn, director, school):
    [badis_october] = [s for s in subscriptions.for_student(conn, director, school.badis)
                       if s.start_date == "2026-10-01"]
    subscriptions.cancel(conn, director, badis_october.id, "Ne vient plus")
    table = reports.revenue_by_month(conn, director, school.year)
    assert table.rows[1][1:4] == (13000, 13000, 0)


def test_money_reports_are_for_the_director(conn, desk, school):
    with pytest.raises(PermissionDenied):
        reports.revenue_by_month(conn, desk, school.year)
    with pytest.raises(PermissionDenied):
        reports.revenue_by_group(conn, desk, school.year)
    with pytest.raises(PermissionDenied):
        reports.teachers_summary(conn, desk, school.year)


# ----- debt by age -----

def test_debtors_by_age(conn, desk, school):
    table = reports.debtors_by_age(conn, desk)
    assert [c.label for c in table.columns] == [
        "Élève", "N°", "0 à 30 j", "31 à 60 j", "61 à 90 j", "Plus de 90 j", "Total"]
    assert table.rows == (("Cherif Badis", "2026-0002", 4000, 3000, 0, 0, 7000),)
    assert table.row_ids == (school.badis,)
    assert table.total == ("Total", None, 4000, 3000, 0, 0, 7000)
    assert table.period == "Au 18/10/2026"
    # The same total as the debtors list.
    assert table.total[-1] == sum(d.owed_now for d in payments_debtors(conn, desk))

    assert reports.debtors_by_age(conn, desk, group_id=school.a).rows == ()


def payments_debtors(conn, ctx):
    from tadris.services import debt
    return debt.debtors(conn, ctx)


def test_the_credit_pays_the_oldest_late_charges_first():
    charges = (
        ChargeDue(id=1, kind="subscription", description="", amount=3000,
                  due_date="2026-06-01", paid=0, group_id=1),  # 139 days late
        ChargeDue(id=2, kind="subscription", description="", amount=3000,
                  due_date="2026-08-01", paid=1000, group_id=1),  # 78 days
        ChargeDue(id=3, kind="subscription", description="", amount=3000,
                  due_date="2026-10-10", paid=0, group_id=1),  # 8 days
        ChargeDue(id=4, kind="subscription", description="", amount=3000,
                  due_date="2026-11-01", paid=0, group_id=1),  # not due yet
    )
    owed = Debt(student_id=1, as_of="2026-10-18", charges=charges, credit=3500)
    # 3 000 credit clears June, 500 goes on August (2 000 left -> 1 500).
    assert reports.age_buckets(owed, TODAY) == [3000, 0, 1500, 0]
    assert sum(reports.age_buckets(owed, TODAY)) == owed.owed_now


# ----- attendance -----

@pytest.fixture
def marks(conn, director, school):
    """A: 06/10 Amine present, Chahinez absent; 13/10 Amine late, Chahinez excused;
    15/10 cancelled. B: 07/10 Badis present."""
    conn.execute("INSERT INTO rooms(name, capacity) VALUES ('Salle 1', 20)")
    room = conn.execute("SELECT id FROM rooms").fetchone()[0]

    def session(group, teacher, day, status="held"):
        reason = "Enseignant absent" if status == "cancelled" else None
        return conn.execute(
            "INSERT INTO sessions(group_id, room_id, teacher_id, date, start_time, end_time, "
            "status, source, cancellation_reason) VALUES (?, ?, ?, ?, '17:00', '18:30', ?, "
            "'manual', ?)", (group, room, teacher, day, status, reason)).lastrowid

    def enrollment(student):
        return conn.execute("SELECT id FROM enrollments WHERE student_id = ?",
                            (student,)).fetchone()[0]

    def mark(session_id, student, status):
        conn.execute("INSERT INTO attendances(session_id, enrollment_id, status, entered_by) "
                     "VALUES (?, ?, ?, ?)",
                     (session_id, enrollment(student), status, director.user_id))

    first = session(school.a, school.karim, "2026-10-06")
    mark(first, school.amine, "present")
    mark(first, school.chahinez, "absent")
    second = session(school.a, school.karim, "2026-10-13")
    mark(second, school.amine, "late")
    mark(second, school.chahinez, "excused")
    session(school.a, school.karim, "2026-10-15", "cancelled")
    mark(session(school.b, school.samia, "2026-10-07"), school.badis, "present")


def test_attendance_by_group(conn, desk, marks):
    table = reports.attendance_by_group(conn, desk, date(2026, 10, 1), TODAY)
    assert [c.label for c in table.columns] == [
        "Groupe", "Séances", "Présents", "Retards", "Absents", "Excusés", "Taux de présence"]
    assert table.rows == (
        ("A · Maths 4AM", 2, 1, 1, 1, 1, 6666),  # 2 of 3, the excused one left out
        ("B · Physique 4AM", 1, 1, 0, 0, 0, 10000),
    )
    assert table.total == ("Total", 3, 2, 1, 1, 1, 7500)
    assert table.period == "Du 01/10/2026 au 18/10/2026"
    assert cells.display("percent", 6666) == "66,66 %"

    week = reports.attendance_by_group(conn, desk, date(2026, 10, 7), date(2026, 10, 13))
    assert [row[:2] for row in week.rows] == [("A · Maths 4AM", 1), ("B · Physique 4AM", 1)]


def test_attendance_by_student_lowest_rate_first(conn, desk, school, marks):
    table = reports.attendance_by_student(conn, desk, date(2026, 10, 1), TODAY)
    assert table.rows == (
        ("Djebbar Chahinez", "2026-0003", "A · Maths 4AM", 0, 0, 1, 1, 0),
        ("Benali Amine", "2026-0001", "A · Maths 4AM", 1, 1, 0, 0, 10000),
        ("Cherif Badis", "2026-0002", "B · Physique 4AM", 1, 0, 0, 0, 10000),
    )
    assert table.row_ids == (school.chahinez, school.amine, school.badis)
    only_b = reports.attendance_by_student(conn, desk, date(2026, 10, 1), TODAY, school.b)
    assert [row[0] for row in only_b.rows] == ["Cherif Badis"]


def test_an_excused_only_student_has_no_rate(conn, desk, school, marks):
    table = reports.attendance_by_student(conn, desk, date(2026, 10, 13), date(2026, 10, 13))
    assert table.rows[-1][0] == "Djebbar Chahinez" and table.rows[-1][-1] is None


def test_dates_in_the_wrong_order(conn, desk):
    with pytest.raises(ValidationError) as raised:
        reports.attendance_by_group(conn, desk, TODAY, date(2026, 10, 1))
    assert raised.value.message == t["report.dates_order"]


# ----- teachers -----

def test_teachers_summary(conn, director, school):
    table = reports.teachers_summary(conn, director, school.year)
    assert [c.label for c in table.columns] == [
        "Enseignant", "Groupes", "Élèves", "Payé", "En attente", "Mois à préparer",
        "À préparer", "Reste à payer"]
    assert table.rows == (
        # Paid September 3 600; October: Amine's and Chahinez's months, 12 000 x 60 %.
        ("Karim Haddad", 1, 2, 3600, 0, 2, 7200, 7200),
        # September's 2 000 waits in a draft; October's 4 000 x 50 %.
        ("Samia Rahmani", 1, 1, 0, 2000, 1, 2000, 4000),
    )
    assert table.total == ("Total", 2, 3, 3600, 2000, 3, 9200, 11200)
    # À préparer is what a draft built today would pay.
    draft = payouts.build_draft(conn, director, school.karim, TODAY)
    assert payouts.get(conn, director, draft).amount_due == 7200


# ----- students -----

def test_student_list(conn, desk, school):
    table = reports.student_list(conn, desk)
    assert [row[:3] for row in table.rows] == [
        ("2026-0001", "Benali", "Amine"), ("2026-0002", "Cherif", "Badis"),
        ("2026-0003", "Djebbar", "Chahinez")]
    level, groups_col, owed = (column(table, label) for label in ("Niveau", "Groupes",
                                                                  "En retard"))
    assert [row[level] for row in table.rows] == ["4AM"] * 3
    assert table.rows[1][groups_col] == "B · Physique 4AM"
    assert [row[owed] for row in table.rows] == [0, 7000, 0]
    assert table.total[0] == "3 élève(s)" and table.total[owed] == 7000
    assert table.row_ids == (school.amine, school.badis, school.chahinez)

    in_a = reports.student_list(conn, desk, group_id=school.a)
    assert [row[1] for row in in_a.rows] == ["Benali", "Djebbar"]
    assert reports.student_list(conn, desk, status="inactive").rows == ()
    with pytest.raises(ValidationError):
        reports.student_list(conn, desk, status="gone")


# ----- dashboard -----

def test_dashboard_for_the_director(conn, director, school, monkeypatch):
    monkeypatch.setattr(reports, "_now_time", lambda: "10:00")
    board = reports.dashboard(conn, director)
    assert (board.debt_total, board.debtors) == (7000, 1)
    assert [d.student_id for d in board.top_debtors] == [school.badis]
    assert (board.charged, board.settled, board.collected) == (17000, 13000, 14000)
    assert [(g.title, g.enrolled, g.seats, g.rate_bp) for g in board.groups] == [
        ("A", 2, 20, 1000), ("B", 1, 20, 500)]
    assert (board.enrolled, board.seats) == (3, 40)


def test_the_receptionist_sees_no_money_on_the_dashboard(conn, desk, school):
    board = reports.dashboard(conn, desk)
    assert board.charged is None and board.collected is None
    assert board.debt_total == 7000


def test_todays_sessions_left_and_cancelled(conn, director, school, monkeypatch):
    conn.execute("INSERT INTO rooms(name, capacity) VALUES ('Salle 1', 20)")
    for start, end, status in (("09:00", "10:30", "held"), ("17:00", "18:30", "planned"),
                               ("19:00", "20:30", "cancelled")):
        conn.execute(
            "INSERT INTO sessions(group_id, room_id, teacher_id, date, start_time, end_time, "
            "status, source, cancellation_reason) VALUES (?, 1, ?, '2026-10-18', ?, ?, ?, "
            "'manual', ?)", (school.a, school.karim, start, end, status,
                             "Absent" if status == "cancelled" else None))
    monkeypatch.setattr(reports, "_now_time", lambda: "12:00")
    board = reports.dashboard(conn, director)
    assert [s.start_time for s in board.sessions] == ["09:00", "17:00", "19:00"]
    assert board.sessions_held == 2
    assert [s.start_time for s in board.sessions_left] == ["17:00"]


def test_dashboard_money_by_day(conn, director, school):
    """October to the 18th: 6 000 on the 3rd, 8 000 on the 6th. A day is open when it had a
    session (a cancelled one too) or a payment."""
    conn.execute("INSERT INTO rooms(name, capacity) VALUES ('Salle 1', 20)")
    conn.execute(
        "INSERT INTO sessions(group_id, room_id, teacher_id, date, start_time, end_time, "
        "status, source, cancellation_reason) VALUES (?, 1, ?, '2026-10-12', '17:00', "
        "'18:30', 'cancelled', 'manual', 'Absent')", (school.a, school.karim))
    board = reports.dashboard(conn, director)
    assert [d.day for d in board.daily] == [f"2026-10-{d:02d}" for d in range(1, 19)]
    assert sum(d.amount for d in board.daily) == board.collected == 14000
    assert [(d.day[8:], d.amount) for d in board.daily if d.open] == [
        ("03", 6000), ("06", 8000), ("12", 0)]


def test_dashboard_closed_weekdays(conn, director, school):
    assert reports.dashboard(conn, director).closed_weekdays == ()  # no timetable yet
    conn.execute("INSERT INTO rooms(name, capacity) VALUES ('Salle 1', 20)")
    for weekday in (1, 3):
        conn.execute(
            "INSERT INTO timetable_slots(group_id, room_id, weekday, start_time, end_time, "
            "valid_from) VALUES (?, 1, ?, '17:00', '18:30', '2026-09-01')", (school.a, weekday))
    conn.execute(
        "INSERT INTO timetable_slots(group_id, room_id, weekday, start_time, end_time, "
        "valid_from, valid_to) VALUES (?, 1, 5, '17:00', '18:30', '2026-09-01', '2026-10-01')",
        (school.b,))  # Fridays ended: the school is closed then now
    assert reports.dashboard(conn, director).closed_weekdays == (2, 4, 5, 6, 7)


def test_dashboard_backup_of_the_day(conn, director, school, desk):
    def copy(status: str, day: str = "2026-10-18") -> None:
        conn.execute("INSERT INTO backups(file_path, type, status, created_at) "
                     "VALUES ('x.db', 'auto', ?, ?)", (status, f"{day}T09:00:00"))

    assert reports.dashboard(conn, director).backup == "none"
    copy("verified", "2026-10-17")  # yesterday's does not count
    assert reports.dashboard(conn, director).backup == "none"
    copy("failed")
    assert reports.dashboard(conn, director).backup == "failed"
    copy("verified")
    copy("failed")  # one verified copy today is enough
    assert reports.dashboard(conn, director).backup == "verified"
    assert reports.dashboard(conn, desk).backup is None


def test_dashboard_for_the_receptionist_has_no_daily_money(conn, desk, school):
    board = reports.dashboard(conn, desk)
    assert board.daily is None
    assert [d.student_id for d in board.oldest_debtors] == [school.badis]
    assert (board.opens_at, board.closes_at) == ("08:00", "20:00")


# ----- audit log -----

def test_audit_log_filters(conn, director, desk, school):
    first, last = date(2026, 1, 1), date(2027, 12, 31)
    every = reports.audit_log(conn, director, first, last)
    assert every and every[0].id > every[-1].id  # latest first

    paid = reports.audit_log(conn, director, first, last, family="payment")
    assert {e.action for e in paid} == {"payment.record", "payment.void"}
    assert len([e for e in paid if e.action == "payment.record"]) == 5

    voids = reports.audit_log(conn, director, first, last, family="payment.void")
    assert [e.action for e in voids] == ["payment.void"]

    assert reports.audit_log(conn, director, first, last, user_id=desk.user_id) == [
        e for e in every if e.user == "accueil"]

    student = reports.audit_log(conn, director, first, last, family="student",
                                entity_id=school.badis)
    assert {e.entity_id for e in student} == {school.badis}
    assert "payment" in reports.audit_families(conn, director)
    assert reports.family_label("payment") == "Paiements"
    assert reports.family_label("unknown") == "unknown"


def test_audit_log_by_day(conn, director, school):
    today = date.today()  # rows are stamped by SQLite's clock
    assert reports.audit_log(conn, director, today, today)
    assert reports.audit_log(conn, director, date(2020, 1, 1), date(2020, 1, 2)) == []


def test_audit_log_is_for_the_director(conn, desk):
    with pytest.raises(PermissionDenied):
        reports.audit_log(conn, desk, TODAY, TODAY)
    with pytest.raises(PermissionDenied):
        reports.audit_families(conn, desk)


def test_the_journal_reads_in_french(conn, director, school):
    voided = reports.audit_log(conn, director, date(2020, 1, 1), date(2030, 1, 1),
                               family="payment.void")
    table = reports.audit_table(conn, voided, TODAY, TODAY)
    assert [c.label for c in table.columns] == ["Date", "Utilisateur", "Action", "Élément",
                                                "Détail"]
    _when, user, action, item, detail = table.rows[0]
    assert (user, action) == ("directeur", "Paiement annulé")
    assert item.startswith("Paiement n° ")
    assert "Montant : 500 DA" in detail and "Motif : Erreur de saisie" in detail
    assert "Annulé : Oui" in detail
    assert "{" not in detail and "_" not in detail  # no JSON, no internal names left


def test_the_journal_names_what_it_points_to(conn, director, school):
    entries = reports.audit_log(conn, director, date(2020, 1, 1), date(2030, 1, 1),
                                family="enrollment.create")
    table = reports.audit_table(conn, entries, TODAY, TODAY)
    details = [row[4] for row in table.rows]
    assert any("Élève : Benali Amine" in d and "Groupe : A" in d for d in details)
    assert all("Entrée : " in d for d in details)  # the date, as dd/mm/yyyy
    assert any("01/09/2026" in d for d in details)


def automatic_entry(conn) -> None:
    """What the app writes by itself, e.g. a backup."""
    from tadris.db.connection import transaction
    from tadris.services import audit

    with transaction(conn):
        audit.log_system(conn, "backup.create", "backups",
                         new={"id": 1, "file": "tadris.db", "type": "auto"})


def test_the_journal_shows_what_people_did(conn, director, school):
    """The app's own work (backups, sessions prepared, months ended) is not shown, nor
    offered as a type of operation; a failed backup is."""
    automatic_entry(conn)
    shown = reports.audit_log(conn, director, date(2020, 1, 1), date(2030, 1, 1))
    assert shown and all(e.user is not None for e in shown)
    assert "backup" not in reports.audit_families(conn, director)

    from tadris.db.connection import transaction
    from tadris.services import audit
    with transaction(conn):
        audit.log_system(conn, "backup.failed", "backups",
                         new={"id": 2, "file": "x.db", "error": "disk full"})
    failed = reports.audit_log(conn, director, date(2020, 1, 1), date(2030, 1, 1),
                               family="backup")
    assert [(e.action, e.user) for e in failed] == [("backup.failed", None)]
    table = reports.audit_table(conn, failed, TODAY, TODAY)
    assert table.rows[0][1:3] == ("Système", "Sauvegarde échouée")
    assert "backup" in reports.audit_families(conn, director)


def test_audit_fields_show_changes(conn, director, school):
    from tadris.services.journal import Field

    approved = reports.audit_log(conn, director, date(2020, 1, 1), date(2030, 1, 1),
                                 family="payout.approve")[0]
    assert reports.audit_fields(conn, director, approved) == [
        Field("Statut", "Approuvé", before="Brouillon"),
        Field("À payer", "3 600 DA")]


def test_audit_fields_do_not_repeat_the_item(conn, director, desk, school):
    created = reports.audit_log(conn, director, date(2020, 1, 1), date(2030, 1, 1),
                                family="user.create")[0]
    # « Utilisateur : accueil » is the item already: only the role is left.
    assert [(f.label, f.value) for f in reports.audit_fields(conn, director, created)] == [
        ("Rôle", "Réceptionniste")]


# ----- exports -----

def test_excel_export_holds_the_screen_totals(conn, director, school, tmp_path):
    table = reports.revenue_by_month(conn, director, school.year)
    path = printing.report_xlsx(conn, director, table, tmp_path)
    assert path.suffix == ".xlsx" and path.name.startswith("rapport_recettes-par-mois_")

    sheet = load_workbook(path).active
    assert sheet["A1"].value == "Recettes par mois"
    assert sheet["A2"].value == "École Ennour · Année 2026-2027"
    header = [cell.value for cell in sheet[4]]
    assert header == [c.label for c in table.columns]
    body = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=5, max_row=14)]
    assert [tuple(row) for row in body] == list(table.rows)
    total = [cell.value for cell in sheet[15]]
    assert tuple(total) == table.total
    # Excel's own sum of the column is the total.
    assert sum(row[1] for row in body) == total[1] == 29000
    assert sheet["B5"].number_format == '#,##0 "DA"'

    entry = conn.execute("SELECT new_value FROM audit_log WHERE action = 'report.export'"
                         ).fetchone()
    assert json.loads(entry[0]) == {"report": "recettes-par-mois", "format": "xlsx",
                                    "rows": 10, "period": "Année 2026-2027"}


def test_excel_keeps_rates_and_dates_as_values(conn, desk, marks, tmp_path):
    table = reports.attendance_by_group(conn, desk, date(2026, 10, 1), TODAY)
    sheet = load_workbook(printing.report_xlsx(conn, desk, table, tmp_path)).active
    assert sheet["G5"].value == pytest.approx(0.6666)
    assert sheet["G5"].number_format == "0.0%"

    students_table = reports.student_list(conn, desk)
    sheet = load_workbook(printing.report_xlsx(conn, desk, students_table, tmp_path)).active
    assert sheet["D5"].value is None  # no birth date


def test_pdf_export(conn, director, school, tmp_path):
    for table in (reports.revenue_by_month(conn, director, school.year),
                  reports.student_list(conn, director),
                  reports.debtors_by_age(conn, director)):
        path = printing.report_pdf(conn, director, table, tmp_path)
        assert path.suffix == ".pdf" and path.read_bytes().startswith(b"%PDF")
    empty = reports.attendance_by_group(conn, director, TODAY, TODAY)
    assert printing.report_pdf(conn, director, empty, tmp_path).stat().st_size > 0


def test_cells_display():
    assert cells.display("money", 12000) == "12 000 DA"
    assert cells.display("number", 1500) == "1 500"
    assert cells.display("date", "2026-10-18") == "18/10/2026"
    assert cells.display("datetime", "2026-10-18T09:05:00") == "18/10/2026 09:05"
    assert cells.display("text", None) == "—"
    assert cells.display("percent", 10000) == "100 %"


# ----- on the demo seed -----

def test_every_report_agrees_with_the_seed(tmp_path):
    """The demo school: each report runs, and its totals match an independent count."""
    import importlib.util
    from pathlib import Path

    from tadris.db.connection import connect
    from tadris.services import debt

    seed_path = Path(__file__).resolve().parents[1] / "scripts" / "seed_demo.py"
    spec = importlib.util.spec_from_file_location("seed_demo", seed_path)
    seed = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(seed)
    seed.main(tmp_path / "demo.db")
    conn = connect(tmp_path / "demo.db")
    row = conn.execute("SELECT id, username FROM users WHERE role = 'director'").fetchone()
    boss = UserContext(user_id=row[0], username=row[1], role="director")
    year = school_years.current(conn, boss)

    by_month = reports.revenue_by_month(conn, boss, year.id)
    by_group = reports.revenue_by_group(conn, boss, year.id)
    charged, settled = conn.execute(
        "SELECT SUM(c.amount), SUM((SELECT COALESCE(SUM(pa.amount), 0) "
        "FROM payment_allocations pa WHERE pa.charge_id = c.id AND NOT EXISTS "
        "(SELECT 1 FROM payment_voids v WHERE v.payment_id = pa.payment_id))) "
        "FROM charges c WHERE c.status = 'open'").fetchone()
    assert by_month.total[1:3] == by_group.total[2:4] == (charged, settled)
    received = conn.execute(
        "SELECT SUM(amount) FROM payments p WHERE NOT EXISTS "
        "(SELECT 1 FROM payment_voids v WHERE v.payment_id = p.id)").fetchone()[0]
    assert by_month.total[4] == received

    aged = reports.debtors_by_age(conn, boss)
    owing = debt.debtors(conn, boss)
    assert aged.total[-1] == sum(d.owed_now for d in owing)
    assert all(sum(row[2:6]) == row[6] for row in aged.rows)

    summary = reports.teachers_summary(conn, boss, year.id)
    waiting = conn.execute("SELECT SUM(amount_due) FROM teacher_payouts "
                           "WHERE status IN ('draft', 'approved')").fetchone()[0] or 0
    paid = conn.execute("SELECT SUM(amount_paid) FROM teacher_payouts "
                        "WHERE status = 'paid'").fetchone()[0] or 0
    assert summary.total[3:5] == (paid, waiting)

    listed = reports.student_list(conn, boss, status=None)
    assert len(listed.rows) == conn.execute("SELECT COUNT(*) FROM students").fetchone()[0]
    first = date.fromisoformat(year.start_date)
    assert reports.attendance_by_group(conn, boss, first, date.today()).rows  # seed: real dates
    board = reports.dashboard(conn, boss)
    assert board.debt_total == aged.total[-1]
    for table in (by_month, by_group, aged, summary, listed):
        load_workbook(printing.report_xlsx(conn, boss, table, tmp_path))
        printing.report_pdf(conn, boss, table, tmp_path)


# ----- the Journal's wording -----

def test_journal_values():
    from tadris.services import journal

    names = journal.Names({"students": {7: "Benali Amine"}, "class_groups": {3: "Maths 4AM"}})
    assert journal.action_label("payment.record") == "Paiement enregistré"
    assert journal.action_label("phase11.sms") == "phase11.sms"  # unknown: as recorded
    assert journal.entity_label("students", 7, names) == "Élève : Benali Amine"
    assert journal.entity_label("payments", 12, names) == "Paiement n° 12"
    assert journal.entity_label("reports", None, names) == "Rapport"
    new = json.dumps({"student_id": 7, "group_id": 3, "price": 3000, "percentage_bp": 6250,
                      "weekday": 2, "method": "cheque", "is_main": True, "charge_id": 99,
                      "start_date": "2026-09-15", "notes": None, "mystery": "x"})
    assert journal.lines(None, new, names) == [
        "Élève : Benali Amine", "Groupe : Maths 4AM", "Prix : 3 000 DA", "Taux : 62,5 %",
        "Jour : Mardi", "Mode : Chèque", "Contact principal : Oui",
        "Début : 15/09/2026", "mystery : x"]
    marks = json.dumps({"date": "2026-10-18", "changes": [
        {"enrollment_id": 1, "old": None, "new": "present"},
        {"enrollment_id": 2, "old": None, "new": "present"},
        {"enrollment_id": 3, "old": "present", "new": "absent"}]})
    assert journal.lines(None, marks, names) == ["Date : 18/10/2026",
                                                 "Marques : 2 Présent, 1 Absent"]
    export = json.dumps({"report": "liste-eleves", "format": "xlsx", "rows": 3})
    assert journal.summary(None, export, names) == (
        "Rapport : Liste des élèves · Format : Excel · Lignes : 3")
    assert journal.lines('{"key": "school.name", "value": "A"}',
                         '{"key": "school.name", "value": "B"}', names) == [
        "Paramètre : Nom de l'école", "Valeur : A → B"]
    assert len(journal.summary(None, json.dumps({"reason": "x" * 300}), names, 50)) == 50


# ----- exports: layout and safety -----

def _journal_like(rows: int = 3):
    from tadris.domain.models import ReportColumn, ReportTable

    return ReportTable(
        key="journal", title="Journal", period="P",
        columns=(ReportColumn("Date", "datetime"), ReportColumn("Utilisateur"),
                 ReportColumn("Action"), ReportColumn("Élément"), ReportColumn("Détail")),
        rows=tuple(("2026-10-20T21:31:45", "directeur", "Paiement enseignant approuvé",
                    "Paiement enseignant n° 6",
                    "Statut : Brouillon → Approuvé · À payer : 16 500 DA · " * 3)
                   for _ in range(rows)))


def test_pdf_columns_never_overlap_nor_cut_words():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase.pdfmetrics import stringWidth

    from tadris.documents.pdf import FONT, FONT_BOLD, MARGIN, MONO, register_fonts
    from tadris.documents.report_pdf import PAD, _fit, layout

    register_fonts()
    report = _journal_like()
    sizes, wanted = layout(report, 9)
    assert wanted > A4[0] - 2 * MARGIN  # it would rather be landscape
    widths = _fit(sizes, landscape(A4)[0] - 2 * MARGIN)
    assert sum(widths) <= landscape(A4)[0] - 2 * MARGIN + 0.01
    # The date and time fit on their line (they used to run under the next column).
    assert widths[0] >= stringWidth("20/10/2026 21:31", MONO, 9) + 2 * PAD
    # A short text stays whole; a header never breaks a word.
    assert widths[1] >= stringWidth("directeur", FONT, 9) + 2 * PAD
    assert widths[1] >= stringWidth("Utilisateur", FONT_BOLD, 8.5) + 2 * PAD


def test_pdf_text_draws_arrows_and_keeps_amounts_whole():
    from tadris.documents.pdf import register_fonts
    from tadris.documents.report_pdf import markup

    register_fonts()
    assert markup("Brouillon → Approuvé") == (
        'Brouillon <font name="PlexMono">→</font> Approuvé')  # Readex Pro has no arrow
    assert markup("À payer : 16 500 DA") == "À payer : 16 500 DA"
    assert markup("a < b & c") == "a &lt; b &amp; c"


def test_a_wide_report_prints_in_landscape(tmp_path):
    from tadris.documents.pdf import Footer, School
    from tadris.documents.report_pdf import ReportDoc, build_report

    path = build_report(ReportDoc(School("École", "", ""), _journal_like(40), Footer("x")),
                        tmp_path / "j.pdf")
    assert b"/MediaBox [ 0 0 841.8898 595.2756 ]" in path.read_bytes()


def test_excel_text_is_never_a_formula(tmp_path):
    from tadris.documents.report_xlsx import build_report_xlsx
    from tadris.domain.models import ReportColumn, ReportTable

    report = ReportTable(key="x", title="X", period="P",
                         columns=(ReportColumn("Élève"), ReportColumn("Montant", "money")),
                         rows=(('=HYPERLINK("http://x","clic")', 100),))
    sheet = load_workbook(build_report_xlsx(report, "École", tmp_path / "x.xlsx")).active
    assert sheet["A5"].data_type == "s"
    assert sheet["A5"].value == '=HYPERLINK("http://x","clic")'  # shown, never run
    assert sheet["B5"].data_type == "n"


def test_excel_text_columns_fit_their_content(tmp_path):
    from tadris.documents.report_xlsx import TEXT_MAX, build_report_xlsx

    sheet = load_workbook(build_report_xlsx(_journal_like(), "École",
                                            tmp_path / "j.xlsx")).active
    assert sheet.column_dimensions["B"].width == len("Utilisateur") + 2
    assert sheet.column_dimensions["E"].width == TEXT_MAX  # long: capped, and wrapped
    assert sheet["E5"].alignment.wrap_text


def test_the_journal_export_keeps_the_whole_detail(conn, director, school):
    paid = reports.audit_log(conn, director, date(2020, 1, 1), date(2030, 1, 1),
                             family="payment.record")
    table = reports.audit_table(conn, paid, TODAY, TODAY)
    assert all(not row[4].endswith("…") for row in table.rows)
    assert all("Mis en avoir : 0 DA" not in row[4] for row in table.rows)


# ----- the Journal, a page at a time -----

def many_entries(conn, ctx, count: int) -> None:
    from tadris.db.connection import transaction
    from tadris.services import audit

    with transaction(conn):
        for n in range(count):
            audit.log(conn, ctx, "student.update", "students", 1, new={"notes": f"n° {n}"})


def test_the_journal_comes_a_page_at_a_time(conn, director, school):
    many_entries(conn, director, 250)
    first, last = date(2020, 1, 1), date(2030, 1, 1)
    total = reports.audit_count(conn, director, first, last)
    assert total == len(reports.audit_log(conn, director, first, last, limit=10_000)) > 250

    page1 = reports.audit_log(conn, director, first, last)
    assert len(page1) == reports.AUDIT_PAGE == 100
    page2 = reports.audit_log(conn, director, first, last, before_id=page1[-1].id)
    ids = [e.id for e in page1 + page2]
    assert ids == sorted(ids, reverse=True) and len(set(ids)) == 200  # no gap, no repeat
    everything = reports.audit_log(conn, director, first, last, limit=10_000)
    assert [e.id for e in everything[:200]] == ids

    # The filters count the same way the pages read.
    assert reports.audit_count(conn, director, first, last, family="student.update") == 250
    assert reports.audit_count(conn, director, first, last, user_id=10_000) == 0


def test_the_last_day_is_included(conn, director, school):
    today = date.today()  # rows carry SQLite's clock
    assert reports.audit_count(conn, director, today, today) > 0
    assert reports.audit_count(conn, director, today - timedelta(days=3),
                               today - timedelta(days=1)) == 0
