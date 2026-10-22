from datetime import date

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QPoint, Qt
from PySide6.QtWidgets import QLabel, QPushButton, QScrollArea, QWidget

from tadris import config
from tadris.core import dates
from tadris.core.context import UserContext
from tadris.domain.models import Dashboard, DashboardSession, DayCash, Debtor, GroupFill
from tadris.i18n.fr import t
from tadris.services import auth, reports
from tadris.ui.main_window import MainWindow
from tadris.ui.pages import dashboard as dashboard_page
from tadris.ui.pages.dashboard import (
    DashboardPage,
    long_date,
    nice_step,
    share_rows,
    short_date,
)
from tadris.ui.pages.reports import ReportsPage
from tadris.ui.theme import apply_theme, load_fonts
from tadris.ui.widgets import common
from test_reports import TODAY, School, automatic_entry


@pytest.fixture(autouse=True)
def theme(qapp, monkeypatch):
    load_fonts()
    apply_theme(qapp)
    monkeypatch.setattr(reports, "_now_time", lambda: "10:00")


@pytest.fixture(autouse=True)
def documents(monkeypatch, tmp_path):
    opened = []
    monkeypatch.setattr(config, "documents_dir", lambda: tmp_path)
    monkeypatch.setattr(common, "open_document", lambda path: opened.append(path))
    return opened


@pytest.fixture
def director(conn):
    return auth.bootstrap_director(conn, "directeur", "password1", "École Ennour")


@pytest.fixture
def desk(conn, director):
    user_id = auth.create_user(conn, director, "accueil", "password2", "receptionist")
    return UserContext(user_id=user_id, username="accueil", role="receptionist")


@pytest.fixture
def school(conn, director, monkeypatch):
    built = School(conn, director, monkeypatch)
    monkeypatch.setattr(dates, "today", lambda: TODAY)
    return built


def click(qtbot, button):
    qtbot.mouseClick(button, Qt.MouseButton.LeftButton)


def texts(table) -> list[list[str]]:
    return [[table.item(r, c).text() for c in range(table.columnCount())]
            for r in range(table.rowCount())]


def export(qtbot, view, button):
    """Click PDF or Excel and wait for the file: it is written in the background."""
    with qtbot.waitSignal(view.exported, timeout=15000) as done:
        click(qtbot, button)
    return done.args[0]


def double_click_row(table, row: int) -> None:
    table.activated.emit(table.model().index(row, 0))


# ----- dashboard -----

def row_texts(card) -> list[list[str]]:
    """The words of each row shown: labels and buttons, in order."""
    rows = []
    layout = card.rows.layout()
    for i in range(layout.count()):
        row = layout.itemAt(i).widget()
        row.layout().activate()  # placed left to right, even on a page not shown
        words = [w for w in row.findChildren(QWidget) if isinstance(w, (QLabel, QPushButton))
                 and w.text() and not w.isHidden()]
        words.sort(key=lambda w: (w.mapTo(row, QPoint(0, 0)).x(), w.mapTo(row, QPoint(0, 0)).y()))
        rows.append([w.text() for w in words])
    return rows


def first_row(card):
    return card.rows.layout().itemAt(0).widget()


def test_dashboard_texts():
    assert long_date("2026-10-21") == "Mercredi 21 octobre 2026"
    assert short_date("2026-10-02") == "ven. 2 oct."
    assert (nice_step(16500), nice_step(7500), nice_step(0)) == (10000, 5000, 1)


def test_dashboard_for_the_director_on_the_school(qtbot, conn, director, school):
    page = DashboardPage(conn, director)
    qtbot.addWidget(page)
    page.show()
    assert page.day.text() == "Dimanche 18 octobre 2026"
    assert page.hello.text() == "Bonjour, Directeur"
    # October: 6 000 on the 3rd, 8 000 on the 6th; no session yet, so only those days open.
    assert page.cash.title.text() == "Encaissé par jour · octobre"
    assert page.cash.total.text() == "14 000 DA"
    assert (page.cash.picked_day.text(), page.cash.picked_amount.text()) == (
        "dim. 18 oct.", "Fermé")  # today, until the mouse picks a day
    page.cash.chart.hovered.emit(2)
    assert (page.cash.picked_day.text(), page.cash.picked_amount.text()) == (
        "sam. 3 oct.", "6 000 DA")
    assert "1 élève" in page.debtors.subtitle.text() and "7 000 DA" in (
        page.debtors.subtitle.text())
    assert row_texts(page.debtors) == [["Cherif Badis", "7 000 DA", "47 jours",
                                        "Encaisser 7 000 DA"]]
    assert page.debtors.count_note.text().startswith("1 sur 1 · ")
    # The most free seats first: where a new student can go.
    assert row_texts(page.groups) == [["B", "19 libres"],
                                      ["A", "18 libres"]]
    assert page.groups.more.text() == "Les 2 groupes"
    assert page.sessions.empty.isVisibleTo(page)
    assert page.backup.text.text() == "Pas encore de sauvegarde aujourd'hui"


def test_receptionist_sees_no_money_and_no_backup(qtbot, conn, desk, school):
    page = DashboardPage(conn, desk)
    qtbot.addWidget(page)
    assert page.cash is None and page.backup is None
    assert page.hello.text() == "Bonjour, Accueil"
    assert row_texts(page.debtors)[0][0] == "Cherif Badis"


def full_board() -> Dashboard:
    """Nine of everything: more than a dashboard list shows."""
    many = range(9)
    debtors = tuple(Debtor(student_id=i, student_number=f"2026-{i:04d}",
                           student_name=f"Élève {i}", owed_now=1000 * (9 - i), due=1000,
                           credit=0, oldest_overdue="2026-10-01", days_overdue=5 + i)
                    for i in many)
    return Dashboard(
        today="2026-10-18", now="10:00",
        sessions=tuple(DashboardSession(id=i, start_time=f"{8 + i:02d}:00",
                                        end_time=f"{9 + i:02d}:00", group_title=f"G{i}",
                                        room="Salle A", teacher="Karim", status="planned")
                       for i in many),
        ending=(), lapsed=0, debt_total=45000, debtors=12, top_debtors=debtors,
        oldest_debtors=tuple(reversed(debtors)),
        groups=tuple(GroupFill(group_id=i, title=f"G{i}", enrolled=5, seats=10) for i in many),
        charged=0, settled=0, collected=3000,
        daily=tuple(DayCash(day=f"2026-10-{d:02d}", amount=1000 if d % 6 else 0,
                            open=bool(d % 6)) for d in range(1, 19)),
        closed_weekdays=(5,), backup="verified")


def test_rows_are_shared_where_they_are_hidden():
    rows = [40, 40]
    assert share_rows(400, [2, 0], rows) == [2, 0]  # room left over
    assert share_rows(400, [9, 9], rows) == [5, 5]
    assert share_rows(400, [9, 3], rows) == [7, 3]  # the short list shows all, the long one the rest
    assert share_rows(0, [4, 4], rows) == [1, 1]  # always one row, even when nothing fits
    assert share_rows(192, [9, 9], [56, 40]) == [2, 2]  # taller rows take more room


def dashboard_at(qtbot, conn, ctx, width: int, height: int) -> DashboardPage:
    page = DashboardPage(conn, ctx)
    qtbot.addWidget(page)
    page.show_board(full_board())
    page.resize(width, height)
    page.show()
    qtbot.wait(10)
    return page


def lists_bottom(page) -> int:
    """The lowest edge of a card, in the page's coordinates."""
    return max(card.mapTo(page, card.rect().bottomLeft()).y()
               for fixed, lists in page.columns for card in (*fixed, *lists))


def test_rows_fill_the_window_and_never_scroll(qtbot, conn, director):
    """1920 x 1080 at 125 % is 1536 x 864 in Qt's units; the window gets about 1536 x 795.
    The lists take the height there is, and stop at the window's bottom edge."""
    page = dashboard_at(qtbot, conn, director, 1536 - 216, 795 - 24)
    assert page.findChildren(QScrollArea) == []
    assert page.debtors.shown >= 3
    assert page.sessions.shown + page.groups.shown >= 6
    assert lists_bottom(page) <= page.height()
    assert page.debtors.count_note.text().startswith(f"{page.debtors.shown} sur 12 · ")
    assert page.backup.text.text() == "Sauvegarde du jour vérifiée"


def test_a_smaller_window_shows_fewer_rows(qtbot, conn, director):
    big = dashboard_at(qtbot, conn, director, 1320, 900)
    small = dashboard_at(qtbot, conn, director, 1064, 680)
    assert 1 <= small.debtors.shown < big.debtors.shown
    assert small.sessions.shown + small.groups.shown < big.sessions.shown + big.groups.shown
    assert lists_bottom(small) <= small.height()
    assert small.width() == 1064  # long names are cut short, the page is not widened


def test_sorting_the_debts_by_age(qtbot, conn, director):
    page = dashboard_at(qtbot, conn, director, 1320, 900)
    assert row_texts(page.debtors)[0][0] == "Élève 0"  # the largest
    click(qtbot, page.debtors.sort_buttons["age"])
    assert row_texts(page.debtors)[0][:3] == ["Élève 8", "1 000 DA", "13 jours"]  # the oldest


def test_dashboard_opens_a_debtor(qtbot, conn, director, school):
    page = DashboardPage(conn, director)
    qtbot.addWidget(page)
    with qtbot.waitSignal(page.open_student) as opened:
        qtbot.mouseDClick(first_row(page.debtors), Qt.MouseButton.LeftButton)
    assert opened.args == [school.badis]


def test_collect_opens_the_payment_dialog(qtbot, conn, director, school, monkeypatch):
    page = DashboardPage(conn, director)
    qtbot.addWidget(page)
    opened = []

    class FakeDialog:
        def __init__(self, conn, ctx, student, parent=None):
            opened.append(student.id)

    monkeypatch.setattr(dashboard_page, "PaymentDialog", FakeDialog)
    monkeypatch.setattr(dashboard_page, "ask", lambda dialog: False)
    button = next(b for b in first_row(page.debtors).findChildren(QPushButton)
                  if b.text() == "Encaisser 7 000 DA")
    click(qtbot, button)
    assert opened == [school.badis]


def test_top_bar_buttons(qtbot, conn, director, school, monkeypatch):
    page = DashboardPage(conn, director)
    qtbot.addWidget(page)
    page.show()
    assert page.add_student.text() == "Ajouter un élève"
    assert page.collect_payment.text() == "Encaisser un paiement"
    assert page.add_student.isVisible() and page.collect_payment.isVisible()

    class NewStudent:  # « Ajouter un élève »: the new student's page opens
        def __init__(self, conn, ctx, parent=None):
            self.student_id = school.amine

    monkeypatch.setattr(dashboard_page, "StudentDialog", NewStudent)
    monkeypatch.setattr(dashboard_page, "ask", lambda dialog: True)
    with qtbot.waitSignal(page.open_student) as opened:
        click(qtbot, page.add_student)
    assert opened.args == [school.amine]

    paid = []

    class Picker:  # « Encaisser un paiement »: choose the student, then pay
        def __init__(self, conn, ctx, parent=None):
            self.student_id = school.badis

    class Payment:
        def __init__(self, conn, ctx, student, parent=None):
            paid.append(student.id)

    monkeypatch.setattr(dashboard_page, "StudentPicker", Picker)
    monkeypatch.setattr(dashboard_page, "PaymentDialog", Payment)
    monkeypatch.setattr(dashboard_page, "ask", lambda dialog: not isinstance(dialog, Payment))
    click(qtbot, page.collect_payment)
    assert paid == [school.badis]


def test_the_student_picker(qtbot, conn, director, school):
    from tadris.ui.student_picker import StudentPicker

    picker = StudentPicker(conn, director)
    qtbot.addWidget(picker)
    table = picker.table
    # Nothing typed: the late payer first, then everyone else.
    assert table.item(0, 1).text() == "Cherif Badis"
    assert table.item(0, 3).text() == "7 000 DA"
    assert table.rowCount() == 3 and picker.submit.isEnabled()
    picker.search.setText("benali")
    picker._find()
    assert [table.item(r, 1).text() for r in range(table.rowCount())] == ["Benali Amine"]
    assert table.item(0, 3).text() == "—"
    picker.search.setText("zzz")
    picker._find()
    assert table.rowCount() == 0 and not picker.submit.isEnabled()
    assert picker.hint.text() == "Aucun élève actif ne correspond à « zzz »."
    picker.search.setText("")
    picker._find()
    click(qtbot, picker.submit)
    assert picker.student_id == school.badis


def test_links_open_their_page(qtbot, conn, director, school):
    window = MainWindow(conn, director)
    qtbot.addWidget(window)
    window.nav_buttons["dashboard"].click()
    page = window.pages.currentWidget()
    assert isinstance(page, DashboardPage)
    click(qtbot, page.debtors.more)
    assert window.nav_buttons["payments"].isChecked()
    window.back()
    click(qtbot, page.backup.more)
    assert window.nav_buttons["settings"].isChecked()
    settings = window.pages.currentWidget()
    assert settings.tabs.currentWidget() is settings.backups


def test_every_table_header_lines_up_with_its_cells(qtbot, conn, director, school):
    """Every page and tab: a centered date or a right-aligned amount has its header above
    it, not at the left of the column."""
    from PySide6.QtWidgets import QTableWidget, QTabWidget

    window = MainWindow(conn, director)
    qtbot.addWidget(window)
    window.show()
    checked = 0

    def check(page) -> None:
        nonlocal checked
        for table in page.findChildren(QTableWidget):
            if not table.isVisible() or table.rowCount() == 0:
                continue
            for c in range(table.columnCount()):
                cell, head = table.item(0, c), table.horizontalHeaderItem(c)
                if cell is None or head is None:
                    continue
                mask = Qt.AlignmentFlag.AlignHorizontal_Mask
                cell_align = Qt.AlignmentFlag(cell.textAlignment()) & mask
                head_align = Qt.AlignmentFlag(head.textAlignment()) & mask
                assert head_align == cell_align or (
                    cell_align == Qt.AlignmentFlag.AlignLeft and not head_align), (
                    head.text(), cell_align, head_align)
                checked += 1

    for button in window.nav_buttons.values():
        button.click()
        qtbot.wait(1)
        page = window.pages.currentWidget()
        check(page)
        for tabs in page.findChildren(QTabWidget):
            for i in range(tabs.count()):
                tabs.setCurrentIndex(i)
                qtbot.wait(1)
                check(page)
    assert checked > 20  # the school's lists have rows to look at


# ----- reports page -----

def tab_labels(page) -> list[str]:
    return [page.tabs.tabText(i) for i in range(page.tabs.count())]


def test_tabs_follow_the_rights(qtbot, conn, director, desk, school):
    boss = ReportsPage(conn, director)
    front = ReportsPage(conn, desk)
    qtbot.addWidget(boss)
    qtbot.addWidget(front)
    assert tab_labels(boss) == ["Recettes", "Retards", "Présences", "Enseignants", "Élèves",
                                "Journal"]
    assert tab_labels(front) == ["Retards", "Présences", "Élèves"]


def test_revenue_tab_shows_the_service_figures(qtbot, conn, director, school):
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    view = page.views["revenue"]
    shown = texts(view.table)
    assert shown[0] == ["Septembre 2026", "12 000 DA", "9 000 DA", "3 000 DA", "9 000 DA"]
    assert shown[-1] == ["Total", "29 000 DA", "22 000 DA", "7 000 DA", "23 000 DA"]
    assert view.table.item(len(shown) - 1, 0).font().weight() > view.table.item(
        0, 0).font().weight()
    assert view.period.text() == "Année 2026-2027"

    view.view.setCurrentIndex(view.view.findData("group"))
    assert view.month_box.isVisibleTo(view)
    view.month.setCurrentIndex(view.month.findData("2026-10-01"))
    assert view.period.text() == "Octobre 2026"
    assert [row[0] for row in texts(view.table)] == [
        "A · Maths 4AM", "B · Physique 4AM", "Frais d'inscription", "Total"]


def test_excel_button_writes_what_the_screen_shows(qtbot, conn, director, school, documents):
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    view = page.views["revenue"]
    export(qtbot, view, view.excel)
    [path] = documents
    sheet = load_workbook(path).active
    total_row = 4 + len(view.report.rows) + 1
    assert [cell.value for cell in sheet[total_row]] == list(view.report.total)
    assert view.rows_shown()[-1][1] == "29 000 DA" == f"{sheet.cell(total_row, 2).value:,} DA"\
        .replace(",", " ")

    export(qtbot, view, view.pdf)
    assert documents[-1].suffix == ".pdf"


def test_debtors_tab_opens_the_student(qtbot, conn, desk, school):
    page = ReportsPage(conn, desk)
    qtbot.addWidget(page)
    view = page.views["debtors"]
    assert texts(view.table)[0] == ["Cherif Badis", "2026-0002", "4 000 DA", "3 000 DA",
                                    "0 DA", "0 DA", "7 000 DA"]
    with qtbot.waitSignal(page.open_student) as opened:
        double_click_row(view.table, 0)
    assert opened.args == [school.badis]
    # The total row opens nobody.
    with qtbot.assertNotEmitted(page.open_student):
        double_click_row(view.table, 1)


def test_attendance_tab_bad_dates_show_a_banner(qtbot, conn, desk, school):
    page = ReportsPage(conn, desk)
    qtbot.addWidget(page)
    page.tabs.setCurrentIndex(1)
    view = page.views["attendance"]
    view.first.set_value(date(2026, 10, 18))
    view.last.set_value(date(2026, 10, 1))
    assert view.banner.text() == t["report.dates_order"]
    view.last.set_value(date(2026, 10, 31))
    assert view.banner.isHidden()
    assert view.empty.isVisibleTo(view)  # no attendance recorded


def test_student_tab_filters(qtbot, conn, desk, school):
    page = ReportsPage(conn, desk)
    qtbot.addWidget(page)
    page.tabs.setCurrentIndex(2)
    view = page.views["students"]
    assert [row[1] for row in texts(view.table)] == ["Benali", "Cherif", "Djebbar", ""]
    assert texts(view.table)[-1][0] == "3 élève(s)"
    view.group.setCurrentIndex(view.group.findData(school.b))
    assert [row[1] for row in texts(view.table)] == ["Cherif", ""]


def test_audit_tab(qtbot, conn, director, desk, school):
    automatic_entry(conn)
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    page.tabs.setCurrentIndex(5)
    view = page.views["audit"]
    today = date.today()  # audit rows carry SQLite's clock
    view.first.set_value(today)
    view.last.set_value(today)
    assert view.table.rowCount() == len(view.entries) > 0

    assert "Sauvegarde créée" not in {row[2] for row in texts(view.table)}
    assert all(e.user is not None for e in view.entries)  # the app's own work: not shown
    assert view.user.findText("Système") == -1

    view.family.setCurrentIndex(view.family.findData("payment"))
    assert {row[2] for row in texts(view.table)} == {"Paiement enregistré", "Paiement annulé"}
    view.table.selectRow(0)
    detail = view.detail.toPlainText()
    assert detail.startswith(texts(view.table)[0][2] + " par ")  # « Paiement … par accueil »
    assert " à " in detail and "Paiement n° " in detail
    assert "Montant" in detail and "{" not in detail

    view.user.setCurrentIndex(view.user.findData(desk.user_id))
    assert view.table.rowCount() == 0 and view.empty.isVisibleTo(view)

    view.user.setCurrentIndex(0)
    view.entity.setText("abc")
    view.entity.editingFinished.emit()
    assert view.banner.text() == t["report.entity_invalid"]


def test_journal_detail_card():
    from tadris.services.journal import Field
    from tadris.ui.pages.reports import detail_html

    html = detail_html("Remboursement approuvé par directeur", "le 18/10/2026 à 10:00",
                       [Field("Statut", "Approuvé", before="En attente"),
                        Field("Motif", "<départ>")])
    assert "<b>Approuvé</b>" in html and "En attente" in html
    assert "&lt;départ&gt;" in html  # what was typed is escaped
    empty = detail_html("Déconnexion par kadam", "le 18/10/2026 à 15:30", [])
    assert "Aucun autre détail enregistré." in empty


def test_redrawing_the_dashboard_many_times_is_safe(qtbot, conn, director):
    """Shrinking and growing the lists, and going to another page and back, again and again:
    a cell reused across redraws once corrupted memory (a crash a few clicks later)."""
    window = MainWindow(conn, director)
    qtbot.addWidget(window)
    window.show()
    dashboard = window.pages.currentWidget()
    for height in (500, 800, 450, 900, 600) * 4:
        dashboard.show_board(full_board())
        window.resize(1300, height)
        qtbot.wait(1)
        window.nav_buttons["students"].click()
        window.nav_buttons["dashboard"].click()
    assert dashboard.sessions.rows.count() == dashboard.sessions.shown


def test_every_tab_exports_pdf_and_excel(qtbot, conn, director, school, documents):
    """Each tab, each view: the PDF and the Excel file are written, open, and hold the rows
    on screen."""
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    variants = [("revenue", None), ("revenue", ("view", "group")), ("debtors", None),
                ("attendance", None), ("attendance", ("view", "student")),
                ("teachers", None), ("students", None), ("audit", None)]
    keys = list(page.views)
    for key, change in variants:
        page.tabs.setCurrentIndex(keys.index(key))
        view = page.views[key]
        if key == "audit":
            view.first.set_value(date(2020, 1, 1))
        if change is not None:
            combo = getattr(view, change[0])
            combo.setCurrentIndex(combo.findData(change[1]))
        export(qtbot, view, view.pdf)
        assert documents[-1].suffix == ".pdf" and documents[-1].read_bytes()[:4] == b"%PDF"
        export(qtbot, view, view.excel)
        sheet = load_workbook(documents[-1]).active
        assert sheet["A1"].value == view.report.title
        body = [row for row in sheet.iter_rows(min_row=5, max_row=4 + len(view.report.rows),
                                               values_only=True)]
        assert len(body) == len(view.report.rows), key
    assert len(documents) == 2 * len(variants)


def test_the_window_answers_while_an_export_is_written(qtbot, conn, director, school,
                                                        documents, monkeypatch):
    """The file is built outside the window's thread: the buttons wait, a line says what is
    being prepared, and the window keeps handling events until the file is there."""
    import threading

    from tadris.services import printing

    release = threading.Event()
    build = printing.ReportJob.build

    def slow_build(job):
        release.wait(5)  # a long PDF
        return build(job)

    monkeypatch.setattr(printing.ReportJob, "build", slow_build)
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    page.show()
    view = page.views["revenue"]
    click(qtbot, view.pdf)
    qtbot.wait(100)  # the event loop runs: the window is not frozen
    assert not view.pdf.isEnabled() and not view.excel.isEnabled()
    assert view.busy.isVisible() and view.busy.text() == "Préparation du PDF…"
    assert documents == []
    with qtbot.waitSignal(view.exported, timeout=5000):
        release.set()
    assert view.pdf.isEnabled() and view.busy.isHidden()
    assert len(documents) == 1
    entries = conn.execute("SELECT COUNT(*) FROM audit_log WHERE action = 'report.export'"
                           ).fetchone()[0]
    assert entries == 1  # recorded once written, on the window's thread


def test_a_failed_export_says_so(qtbot, conn, director, school, documents, monkeypatch):
    from tadris.services import printing

    def broken(job):
        raise OSError("disk full")

    monkeypatch.setattr(printing.ReportJob, "build", broken)
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    view = page.views["revenue"]
    click(qtbot, view.excel)
    qtbot.waitUntil(lambda: view.excel.isEnabled(), timeout=5000)
    assert view.banner.text() == t["print.failed"] or "disk full" in view.banner.text()
    assert documents == []


def test_refilling_a_table_is_fast(qtbot):
    """A table whose columns fit their content, refilled while it already has rows: Qt used
    to measure every column again for each cell (64 rows x 10 columns: 3.5 s, the window
    « not responding »). The columns now fit once, at the end."""
    import time

    from PySide6.QtWidgets import QApplication, QHeaderView

    from tadris.ui.widgets.tables import fill, fit_columns, item, ruled_table

    def rows():
        return [[item(f"cellule {r} {c} " * 3) for c in range(10)] for r in range(64)]

    table = ruled_table([str(c) for c in range(10)])
    qtbot.addWidget(table)
    table.resize(1200, 600)
    table.show()
    fill(table, rows())
    fit_columns(table, stretch=(5,))
    qtbot.wait(20)  # laid out on screen, as when the user comes back to the page
    started = time.perf_counter()
    fill(table, rows())
    QApplication.processEvents()
    assert time.perf_counter() - started < 0.5
    modes = [table.horizontalHeader().sectionResizeMode(c) for c in range(10)]
    assert modes.count(QHeaderView.ResizeMode.ResizeToContents) == 9  # still fitted
    assert table.rowCount() == 64 and table.item(63, 9).text().startswith("cellule 63 9")


def test_the_journal_loads_more_as_it_is_scrolled(qtbot, conn, director, school, documents):
    from test_reports import many_entries

    many_entries(conn, director, 250)
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    page.resize(1300, 700)
    page.show()
    page.tabs.setCurrentIndex(5)
    view = page.views["audit"]
    view.first.set_value(date(2020, 1, 1))
    view.last.set_value(date(2030, 1, 1))
    qtbot.wait(20)
    total = reports.audit_count(conn, director, date(2020, 1, 1), date(2030, 1, 1))
    assert view.table.rowCount() == len(view.entries) == 100
    assert view.note.text().startswith(f"100 sur {total} opérations")

    bar = view.table.verticalScrollBar()

    def at_the_end() -> bool:
        bar.setValue(bar.maximum())  # the end of the list
        return view.table.rowCount() == 200

    qtbot.waitUntil(at_the_end, timeout=3000)
    assert bar.value() > 0  # the list did not jump back to the top
    assert len(view.entries) == len(view.report.rows) == 200
    assert view.note.text().startswith(f"200 sur {total} opérations")

    view.table.selectRow(150)  # a row of the second page: its own detail
    assert view.entries[150].id < view.entries[99].id
    assert view.detail.toPlainText().startswith(view.report.rows[150][2])

    while view.table.rowCount() < total:
        bar.setValue(bar.maximum())
        qtbot.wait(5)
    assert view.table.rowCount() == total and view.note.text() == f"{total} opération(s)."
    bar.setValue(bar.maximum())
    qtbot.wait(20)
    assert view.table.rowCount() == total  # nothing more to load

    path = export(qtbot, view, view.excel)  # the file holds every entry, not the pages
    sheet = load_workbook(path).active
    assert sheet.max_row - 4 == total


def test_a_new_filter_starts_again_at_the_latest(qtbot, conn, director, school):
    from test_reports import many_entries

    many_entries(conn, director, 150)
    page = ReportsPage(conn, director)
    qtbot.addWidget(page)
    page.resize(1300, 700)
    page.show()
    page.tabs.setCurrentIndex(5)
    view = page.views["audit"]
    view.first.set_value(date(2020, 1, 1))
    view.last.set_value(date(2030, 1, 1))  # the entries carry SQLite's clock, not TODAY
    bar = view.table.verticalScrollBar()

    def scrolled_down() -> bool:  # as a user does: down until the next page is there
        bar.setValue(bar.maximum())
        return view.table.rowCount() > 100

    qtbot.waitUntil(scrolled_down, timeout=3000)
    view.family.setCurrentIndex(view.family.findData("payment"))
    assert view.table.rowCount() == len(view.entries) <= 100
    assert {e.action.split(".")[0] for e in view.entries} == {"payment"}
